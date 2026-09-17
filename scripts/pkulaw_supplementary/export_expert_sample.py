"""抽 20 案（分层：短/中/长）导出 xlsx 给专家审脱敏 + 出题质量"""
import json
import random
import sys
from pathlib import Path

REPO = Path("/aifs4su/hansirui/boqin/abuse/DV-JusticeBench")
sys.path.insert(0, str(REPO))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CASE_FILE = REPO / "data/pkulaw_2023_2025_release/pipeline_ready_cases.json"
CLASSIFIED = REPO / "data/pkulaw_filter/stage2_classified.jsonl"
MASKED = REPO / "data/pkulaw_filter/prepare_output/masked_cases.jsonl"
QUESTIONS = REPO / "data/pkulaw_filter/prepare_output/questions.jsonl"
OUT_XLSX = REPO / "data/pkulaw_filter/expert_review_sample_20.xlsx"

SEED = 42
random.seed(SEED)


def main():
    cases = json.load(open(CASE_FILE))
    classification = {json.loads(l)["case_id"]: json.loads(l)
                      for l in open(CLASSIFIED) if l.strip()}
    masked = {json.loads(l)["case_id"]: json.loads(l)
              for l in open(MASKED) if l.strip()}
    questions = {json.loads(l)["case_id"]: json.loads(l)
                 for l in open(QUESTIONS) if l.strip()}

    # 分层抽样：按 case_text 长度分短/中/长
    prepared_ids = list(masked.keys())
    with_len = [(cid, len(cases[cid].get("case_text", ""))) for cid in prepared_ids]
    with_len.sort(key=lambda x: x[1])
    N = len(with_len)
    short = [cid for cid, _ in with_len[: N // 3]]
    mid = [cid for cid, _ in with_len[N // 3: 2 * N // 3]]
    long = [cid for cid, _ in with_len[2 * N // 3:]]

    # 再按 label 分层：主要抽 primary，也带几个 secondary(court=true)
    primary_ids = [cid for cid in prepared_ids if classification[cid]["label"] == "primary"]
    borderline_ids = [cid for cid in prepared_ids if classification[cid]["label"] == "secondary"]

    sample = []
    for bucket, n in [(short, 6), (mid, 7), (long, 5)]:
        # 优先 primary，不够时用 borderline 补
        pri = [c for c in bucket if c in primary_ids]
        bor = [c for c in bucket if c in borderline_ids]
        random.shuffle(pri)
        random.shuffle(bor)
        take = pri[:n]
        if len(take) < n:
            take += bor[:n - len(take)]
        sample.extend(take)

    # 加 2 个 borderline 增加类别覆盖
    extras = [c for c in borderline_ids if c not in sample]
    random.shuffle(extras)
    sample.extend(extras[:2])

    print(f"抽样 {len(sample)} 案")
    for cid in sample:
        lab = classification[cid]["label"]
        ln = len(cases[cid].get("case_text", ""))
        print(f"  · {cid[-10:]} [{lab}] case_text={ln}")

    # 写 xlsx
    wb = Workbook()
    ws = wb.active
    ws.title = "专家审阅"

    HEADER = [
        ("A", "case_id", 30),
        ("B", "分类", 10),
        ("C", "case_text长度", 10),
        ("D", "原始标题", 40),
        ("E", "脱敏标题", 40),
        ("F", "原始案情", 80),
        ("G", "脱敏案情", 80),
        ("H", "原始裁判", 80),
        ("I", "脱敏裁判", 80),
        ("J", "问题1", 55),
        ("K", "问题2", 55),
        ("L", "问题3", 55),
        ("M", "问题4", 55),
        ("N", "问题5", 55),
        ("O", "脱敏质量\n(1-5)", 10),
        ("P", "出题质量\n(1-5)", 10),
        ("Q", "备注（专家填写）", 40),
    ]
    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")
    for col, name, width in HEADER:
        cell = ws[f"{col}1"]
        cell.value = name
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 32

    for row_idx, cid in enumerate(sample, start=2):
        c = cases[cid]
        m = masked[cid]
        q = questions[cid]["questions"]
        cls = classification[cid]

        row = [
            cid,
            cls["label"],
            len(c.get("case_text", "")),
            c.get("title", ""),
            m.get("title_masked", ""),
            c.get("case_text", ""),
            m.get("case_text_masked", ""),
            c.get("judge_decision", ""),
            m.get("judge_decision_masked", ""),
            *[q[i]["question"] if i < len(q) else "" for i in range(5)],
            "", "", "",  # 空的评分列
        ]
        for i, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=i, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_idx].height = 240  # 高一点让长文本展开

    # 冻结首行 + 首列
    ws.freeze_panes = "B2"

    # 第二页：抽样说明
    ws2 = wb.create_sheet("抽样说明")
    ws2["A1"] = "抽样说明"
    ws2["A1"].font = Font(bold=True, size=14)
    notes = [
        "",
        f"总体来源：231 个已完成 prepare 阶段的案例（primary + secondary&court_addressed_dv=true）",
        f"抽样规模：{len(sample)} 案",
        f"抽样方法：按 case_text 长度分为短/中/长三档，每档抽 6/7/5，优先 primary 类；额外加 2 个 secondary 覆盖类别",
        f"随机种子：{SEED}",
        "",
        "评分建议（专家填写 O/P 列）：",
        "  5 分 = 无问题，可直接进入正式实验",
        "  4 分 = 有小瑕疵但不影响实验有效性",
        "  3 分 = 存在明显问题，需要调整 prompt 或流程",
        "  2 分 = 严重问题，需要重跑",
        "  1 分 = 不可用",
        "",
        "关注点：",
        "  脱敏质量：人名/地名/时间/编号是否完整替换？金额是否保留？是否漏敏或误敏？跨字段编号是否一致（某男1 是否始终指同一人）？",
        "  出题质量：问题是否聚焦争议？是否覆盖法律认定/证据/价值判断？是否避免事实性问题？表述是否完整（有无被截断）？",
    ]
    for i, text in enumerate(notes, start=2):
        ws2.cell(row=i, column=1, value=text)
    ws2.column_dimensions["A"].width = 100

    wb.save(OUT_XLSX)
    print(f"\n[导出] {OUT_XLSX}")
    print(f"  ({OUT_XLSX.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
